"""Build a multi-sheet tax computation workbook (openpyxl)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from oe_engine_app.schemas.tax_return_export import (
    InterestLine,
    TaxReturnExportRequest,
)

_THIN = Side(style="thin", color="1F4E79")
_MEDIUM = Side(style="medium", color="1F4E79")
_DOUBLE = Side(style="double", color="000000")
_HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_BOLD = Font(bold=True, name="Calibri", size=11)
_TITLE = Font(bold=True, name="Calibri", size=12, underline="single")
_MONEY = '#,##0.00;(#,##0.00);"-"'
_MONEY_INT = '#,##0;(#,##0);"-"'


def _money(ws, row: int, col: int, value: float | None, *, integer: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=float(value or 0))
    cell.number_format = _MONEY_INT if integer else _MONEY
    cell.alignment = Alignment(horizontal="right")


def _label(ws, row: int, col: int, text: str, *, bold: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    if bold:
        cell.font = _BOLD


def _underline(cell, double: bool = False) -> None:
    side = _DOUBLE if double else Side(style="thin", color="000000")
    cell.border = Border(bottom=side)


def _bracket_amount(value: float) -> float:
    """Store negatives so Excel accounting format shows parentheses."""
    return -abs(float(value or 0))


def _ya_display(raw: str) -> str:
    s = (raw or "").strip()
    if "_" in s and len(s) >= 7:
        a, b = s.split("_", 1)
        return f"{a}/{b}"
    return s or "—"


def _set_widths(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _relief_applied(result: dict[str, Any], *needles: str) -> float:
    total = 0.0
    for line in result.get("relief_lines") or []:
        if not isinstance(line, dict):
            continue
        blob = " ".join(
            str(line.get(k) or "")
            for k in ("compare_group_id", "display_name", "engine_binding_kind")
        ).lower()
        if any(n.lower() in blob for n in needles):
            total += float(line.get("applied") or 0)
    return total


def _interest_wht_total(lines: list[InterestLine]) -> float:
    return sum(float(line.wht or 0) for line in lines)


def _resolve_wht_credit(result: dict[str, Any], interest_lines: list[InterestLine]) -> float:
    """WHT credit for TAX CALCULATION.

    Prefer engine ``wht_credit``. If the calculate snapshot omitted WHT but the
    profile interest schedule (SAVINGS/FD sheets) has amounts, use those —
    capped after APIT the same way as calculate.
    """
    engine_credit = float(result.get("wht_credit") or 0)
    if engine_credit > 0:
        return engine_credit

    claimed = float(result.get("wht_already_paid") or 0)
    if claimed <= 0:
        claimed = _interest_wht_total(interest_lines)
    if claimed <= 0:
        return 0.0

    tax = float(result.get("tax_payable") or 0)
    apit = float(result.get("apit_credit") or result.get("apit_already_paid") or 0)
    remaining = max(0.0, tax - apit)
    return min(claimed, remaining)


def _build_tax_calculation(wb: Workbook, body: TaxReturnExportRequest) -> None:
    ws = wb.active
    ws.title = "TAX CALCULATION"
    header = body.header
    income = body.income_breakdown
    result = body.calculate_result

    ya = header.assessment_year_label or _ya_display(str(result.get("assessment_year") or ""))
    name = header.full_name or "Taxpayer"
    tin = header.tin or ""

    ws["B2"] = name.upper()
    ws["B2"].font = _TITLE
    _label(ws, 3, 2, "TIN Number")
    ws["C3"] = tin
    ws["B4"] = f"Tax Computation for the Y/A {ya}"
    ws["B4"].font = _TITLE
    if header.reference_code:
        ws["F4"] = header.reference_code

    employment = float(result.get("employment_income") or income.employment or 0)
    interest = float(income.interest or 0)
    if interest <= 0:
        interest = float(result.get("investment_income") or income.investment or 0)
    business = float(result.get("business_income") or income.business or 0)
    rents = float(income.rents or 0)
    other = float(result.get("other_income") or income.other or 0)
    assessable = float(result.get("gross_income") or 0)

    row = 7
    _label(ws, row, 2, "Employment Income")
    _money(ws, row, 6, employment, integer=True)
    row += 1
    _label(ws, row, 2, "Interest Income")
    _money(ws, row, 6, interest, integer=True)
    if business > 0:
        row += 1
        _label(ws, row, 2, "Business Income")
        _money(ws, row, 6, business, integer=True)
    if rents > 0:
        row += 1
        _label(ws, row, 2, "Rental Income")
        _money(ws, row, 6, rents, integer=True)
    if other > 0:
        row += 1
        _label(ws, row, 2, "Other Income")
        _money(ws, row, 6, other, integer=True)
    row += 1
    _label(ws, row, 2, "Assessable Income", bold=True)
    _money(ws, row, 6, assessable, integer=True)
    _underline(ws.cell(row=row, column=6))

    row += 2
    # Section header only — no amount (avoids a dash/total that double-counts detail lines).
    _label(ws, row, 2, "Less- Sec -52", bold=True)

    # Applied Sec 52 / Fifth Schedule reliefs listed under the header.
    row += 2
    personal = _relief_applied(result, "personal_relief", "tax free", "personal relief")
    _label(ws, row, 2, "Tax Free Allowance")
    _money(ws, row, 6, _bracket_amount(personal) if personal else 0, integer=True)

    row += 1
    solar = _relief_applied(result, "solar_panel", "solar")
    _label(ws, row, 2, "Solar Relief")
    _money(ws, row, 6, _bracket_amount(solar) if solar else 0, integer=True)

    # Other applied reliefs not already shown (donations, rental, QP, etc.)
    shown = {"personal_relief", "solar_panel", "sec_52"}
    for line in result.get("relief_lines") or []:
        if not isinstance(line, dict):
            continue
        applied = float(line.get("applied") or 0)
        if applied <= 0:
            continue
        group = str(line.get("compare_group_id") or "").lower()
        kind = str(line.get("engine_binding_kind") or "").lower()
        if any(s in group or s in kind for s in ("personal", "solar", "sec_52", "sec-52")):
            continue
        if group in shown:
            continue
        row += 1
        _label(ws, row, 2, str(line.get("display_name") or group or "Relief"))
        _money(ws, row, 6, _bracket_amount(applied), integer=True)

    row += 2
    taxable = float(result.get("taxable_income") or 0)
    _label(ws, row, 2, "Taxable Income", bold=True)
    _money(ws, row, 6, taxable, integer=True)
    _underline(ws.cell(row=row, column=6), double=True)

    row += 2
    _label(ws, row, 2, "Gross Tax Liability", bold=True)
    row += 1
    for slab in result.get("slab_lines") or []:
        if not isinstance(slab, dict):
            continue
        slice_amt = float(slab.get("slice") or 0)
        if slice_amt <= 0:
            continue
        rate = float(slab.get("rate_percent") or 0)
        tax = float(slab.get("tax") or 0)
        label = str(slab.get("band_label") or f"Band {slab.get('band_index', '')}")
        _label(ws, row, 2, label)
        ws.cell(row=row, column=3, value=slice_amt).number_format = _MONEY_INT
        ws.cell(row=row, column=4, value=rate / 100 if rate > 1 else rate).number_format = "0%"
        _money(ws, row, 5, tax)
        row += 1

    gross_tax = float(result.get("tax_payable") or 0)
    _money(ws, row, 6, gross_tax, integer=True)
    _underline(ws.cell(row=row, column=6), double=True)
    _label(ws, row, 2, "Total", bold=True)
    _money(ws, row, 3, taxable, integer=True)
    _money(ws, row, 5, gross_tax)

    row += 2
    _label(ws, row, 2, "Less - Tax Credits", bold=True)
    row += 1
    wht = _resolve_wht_credit(result, body.interest_lines)
    _label(ws, row, 2, "WHT")
    _money(ws, row, 6, _bracket_amount(wht) if wht else 0, integer=True)
    apit = float(result.get("apit_credit") or result.get("apit_already_paid") or 0)
    if apit > 0:
        row += 1
        _label(ws, row, 2, "APIT")
        _money(ws, row, 6, _bracket_amount(apit), integer=True)

    row += 2
    engine_wht = float(result.get("wht_credit") or 0)
    if wht > engine_wht:
        # Interest-schedule fallback — recompute balance so credits stay consistent.
        final = round(gross_tax - apit - wht, 2)
    else:
        balance = float(result.get("balance_payable") or 0)
        refund = float(result.get("tax_refund") or 0)
        final = -refund if refund > 0 and balance <= 0 else balance
    _label(ws, row, 2, "Balance Tax Payable/(Refund)", bold=True)
    _money(ws, row, 6, final, integer=True)
    _underline(ws.cell(row=row, column=6), double=True)

    _set_widths(ws, {2: 36, 3: 14, 4: 10, 5: 14, 6: 16})


def _interest_headers(ws, row: int = 2, *, upload: bool = False) -> None:
    if upload:
        headers = [
            "Bank",
            "Withholding Agent TIN",
            "WHT Certificate No.",
            "Date",
            "Acc. No./ FD No.",
            "TYPE",
            "FD Amount",
            "Gross Interest",
            "WHT",
        ]
    else:
        headers = [
            "Bank",
            "Source /Type",
            "Withholding Agent TIN",
            "WHT Certificate No.",
            "Date",
            "TYPE",
            "Gross Interest",
            "WHT",
        ]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL


def _write_interest_sheet(
    wb: Workbook,
    title: str,
    lines: list[InterestLine],
    *,
    type_filter: str | None = None,
) -> None:
    ws = wb.create_sheet(title)
    _interest_headers(ws, 2, upload=False)
    filtered = [
        line
        for line in lines
        if type_filter is None or line.type.upper() == type_filter.upper()
    ]
    row = 3
    gross_total = 0.0
    wht_total = 0.0
    for line in filtered:
        ws.cell(row=row, column=1, value=line.bank)
        ws.cell(row=row, column=2, value=line.source_type)
        ws.cell(row=row, column=3, value=line.withholding_agent_tin or "")
        ws.cell(row=row, column=4, value=line.wht_certificate_no or "")
        ws.cell(row=row, column=5, value=line.date or "")
        ws.cell(row=row, column=6, value=line.type)
        _money(ws, row, 7, line.gross_interest)
        _money(ws, row, 8, line.wht)
        gross_total += float(line.gross_interest or 0)
        wht_total += float(line.wht or 0)
        row += 1

    if filtered:
        ws.cell(row=row, column=1, value="Total").font = _BOLD
        ws.cell(row=row, column=1).fill = _TOTAL_FILL
        _money(ws, row, 7, gross_total)
        _money(ws, row, 8, wht_total)
        _underline(ws.cell(row=row, column=7), double=True)
        _underline(ws.cell(row=row, column=8), double=True)

    _set_widths(ws, {1: 22, 2: 18, 3: 18, 4: 20, 5: 12, 6: 10, 7: 14, 8: 12})
    ws.auto_filter.ref = f"A2:H{max(2, row - 1)}"


def _write_upload_file(wb: Workbook, lines: list[InterestLine]) -> None:
    ws = wb.create_sheet("UPLOAD FILE")
    _interest_headers(ws, 2, upload=True)
    row = 3
    for line in lines:
        ws.cell(row=row, column=1, value=line.bank)
        ws.cell(row=row, column=2, value=line.withholding_agent_tin or "")
        ws.cell(row=row, column=3, value=line.wht_certificate_no or "")
        ws.cell(row=row, column=4, value=line.date or "")
        ws.cell(row=row, column=5, value=line.account_or_fd_no or "")
        ws.cell(row=row, column=6, value=line.type)
        if line.fd_amount is not None:
            _money(ws, row, 7, line.fd_amount)
        _money(ws, row, 8, line.gross_interest)
        _money(ws, row, 9, line.wht)
        row += 1

    if lines:
        gross = sum(float(x.gross_interest or 0) for x in lines)
        wht = sum(float(x.wht or 0) for x in lines)
        ws.cell(row=row, column=1, value="Total").font = _BOLD
        _money(ws, row, 8, gross)
        _money(ws, row, 9, wht)
        _underline(ws.cell(row=row, column=8), double=True)
        _underline(ws.cell(row=row, column=9), double=True)

    _set_widths(ws, {1: 22, 2: 18, 3: 20, 4: 12, 5: 16, 6: 10, 7: 14, 8: 14, 9: 12})
    ws.auto_filter.ref = f"A2:I{max(2, row - 1)}"


def _write_pension(wb: Workbook, body: TaxReturnExportRequest) -> None:
    pension = body.pension
    if pension is None:
        return
    ws = wb.create_sheet("Pension Sche.")
    ws["B2"] = "Pensioner"
    ws["C2"] = pension.pensioner_name or body.header.full_name
    ws["B3"] = "Status"
    ws["C3"] = pension.status_note or ""
    ws["B5"] = "Pension Number"
    ws["C5"] = pension.pension_number or ""

    headers = [
        "Pen. Date",
        "Bpen",
        "CLA",
        "OTA",
        "SPA",
        "DSA",
        "MAA",
        "SUR",
        "TDED",
        "TPEN",
        "PNET",
    ]
    for col, text in enumerate(headers, start=2):
        cell = ws.cell(row=7, column=col, value=text)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL

    row = 8
    tded = tpen = pnet = 0.0
    for line in pension.lines:
        ws.cell(row=row, column=2, value=line.pen_date)
        _money(ws, row, 3, line.bpen)
        _money(ws, row, 4, line.cla)
        _money(ws, row, 5, line.ota)
        _money(ws, row, 6, line.spa)
        _money(ws, row, 7, line.dsa)
        _money(ws, row, 8, line.maa)
        _money(ws, row, 9, line.sur)
        _money(ws, row, 10, line.tded)
        _money(ws, row, 11, line.tpen)
        _money(ws, row, 12, line.pnet)
        tded += float(line.tded or 0)
        tpen += float(line.tpen or 0)
        pnet += float(line.pnet or 0)
        row += 1

    if pension.lines:
        ws.cell(row=row, column=2, value="Total").font = _BOLD
        _money(ws, row, 10, tded)
        _money(ws, row, 11, tpen)
        _money(ws, row, 12, pnet)
        _underline(ws.cell(row=row, column=11), double=True)
        _underline(ws.cell(row=row, column=12), double=True)

    _set_widths(ws, {c: 12 for c in range(2, 13)})


def _write_tx_summary(wb: Workbook, body: TaxReturnExportRequest) -> None:
    summary = body.taxable_summary
    if summary is None:
        return
    ws = wb.create_sheet("TX TAXABLE SUMMARY")
    ws["A1"] = "Transaction Semantic — Taxable Income Summary"
    ws["A1"].font = _TITLE
    ws["A2"] = "Total taxable (LKR)"
    _money(ws, 2, 2, summary.total_taxable_lkr)
    ws["A3"] = "Covered months"
    ws["B3"] = summary.covered_month_count
    ws["A4"] = "Missing months"
    ws["B4"] = summary.missing_month_count

    headers = ["Calendar month", "Class key", "Taxable amount (LKR)", "Transaction count"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=col, value=text)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL

    row = 7
    for line in summary.lines:
        ws.cell(row=row, column=1, value=line.calendar_month)
        ws.cell(row=row, column=2, value=line.class_key)
        _money(ws, row, 3, line.taxable_amount_lkr)
        ws.cell(row=row, column=4, value=line.transaction_count)
        row += 1

    _set_widths(ws, {1: 16, 2: 24, 3: 18, 4: 16})


def build_tax_return_workbook(body: TaxReturnExportRequest) -> bytes:
    """Return .xlsx bytes for the assembled tax return export."""
    wb = Workbook()
    _build_tax_calculation(wb, body)
    _write_interest_sheet(wb, "SAVINGS", body.interest_lines, type_filter="SAVINGS")
    _write_interest_sheet(wb, "FD", body.interest_lines, type_filter="FD")
    _write_upload_file(wb, body.interest_lines)
    _write_pension(wb, body)
    _write_tx_summary(wb, body)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
